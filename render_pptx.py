"""
render_pptx.py
==============
Executive PowerPoint presentation generator for Marsh Pitch Generator.
Produces clean, modern, minimalist 16:9 presentations inspired by
top-tier management consulting & corporate advisory standards.
No boxy cards, no clunky footers — pure editorial typography and whitespace.
"""

import os
import logging
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# Marsh McLennan Executive Palette
COLOR_NAVY = RGBColor(0, 38, 100)       # #002664 - Marsh Deep Navy
COLOR_BLUE = RGBColor(0, 114, 206)      # #0072CE - Marsh Cerulean
COLOR_SKY = RGBColor(0, 163, 224)       # #00A3E0 - Accent Sky
COLOR_DARK = RGBColor(26, 32, 44)       # #1A202C - Primary Dark Text
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B - Secondary Slate Text
COLOR_LINE = RGBColor(226, 232, 240)    # #E2E8F0 - Subtle Divider Line
COLOR_WHITE = RGBColor(255, 255, 255)


def render_pptx(pitch_data: Dict[str, Any], output_path: str = 'pitch_output.pptx') -> str:
    """
    Renders an executive, clean, non-boxy presentation matching the generated pitch.
    """
    try:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6] # completely blank canvas

        company_name = pitch_data.get('company_name', 'Client')
        slides_data: List[Dict] = pitch_data.get('slides', [])
        rec_policy = pitch_data.get('recommended_policy', {})

        # -------------------------------------------------------------
        # SLIDE 1: Title Slide (Minimalist Executive Cover)
        # -------------------------------------------------------------
        title_slide = prs.slides.add_slide(blank_layout)

        # Elegant Left Cerulean Accent Bar
        left_accent = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.08), Inches(3.8))
        left_accent.fill.solid()
        left_accent.fill.fore_color.rgb = COLOR_BLUE
        left_accent.line.fill.background()

        # Title & Subtitle Box (Clean typography, no gray cards)
        tbox = title_slide.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(10.5), Inches(4.2))
        tf = tbox.text_frame
        tf.word_wrap = True

        p_pre = tf.paragraphs[0]
        p_pre.text = "MARSH MCLENNAN  |  CORPORATE HEALTH & BENEFITS ADVISORY"
        p_pre.font.name = "Arial"
        p_pre.font.size = Pt(11)
        p_pre.font.bold = True
        p_pre.font.color.rgb = COLOR_BLUE
        p_pre.space_after = Pt(14)

        p1 = tf.add_paragraph()
        p1.text = "Strategic Health & Insurance Pitch"
        p1.font.name = "Arial"
        p1.font.size = Pt(42)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_NAVY
        p1.space_after = Pt(16)

        p2 = tf.add_paragraph()
        p2.text = f"Custom Risk Mitigation & Employee Benefits Proposal for {company_name}"
        p2.font.name = "Arial"
        p2.font.size = Pt(20)
        p2.font.color.rgb = COLOR_MUTED

        # -------------------------------------------------------------
        # SLIDES 2..N: Content Slides (Clean Editorial Layout, No Boxy Cards)
        # -------------------------------------------------------------
        for slide_data in slides_data:
            slide = prs.slides.add_slide(blank_layout)
            title_text = slide_data.get('title', 'Strategic Overview')
            bullets = slide_data.get('bullets', [])

            # Elegant Top Navy Header Bar
            header_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
            header_rect.fill.solid()
            header_rect.fill.fore_color.rgb = COLOR_NAVY
            header_rect.line.fill.background()

            # Thin Accent Line
            accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.04))
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = COLOR_BLUE
            accent_bar.line.fill.background()

            # Slide Title Text
            tbox = slide.shapes.add_textbox(Inches(0.9), Inches(0.2), Inches(11.5), Inches(0.75))
            ttf = tbox.text_frame
            tp = ttf.paragraphs[0]
            tp.text = title_text
            tp.font.name = "Arial"
            tp.font.size = Pt(26)
            tp.font.bold = True
            tp.font.color.rgb = COLOR_WHITE

            # Editorial Bullets (Clean spacing, subtle numbers, no boxy containers)
            num_bullets = max(len(bullets), 1)
            start_y = 1.6
            avail_h = 5.2
            spacing = avail_h / num_bullets

            for idx, bullet in enumerate(bullets):
                current_y = start_y + (idx * spacing)

                if isinstance(bullet, dict):
                    b_text = bullet.get('text', '')
                    source_id = bullet.get('source_chunk_id')
                else:
                    b_text = str(bullet)
                    source_id = None

                # Clean Numeric Indicator
                num_box = slide.shapes.add_textbox(Inches(0.9), Inches(current_y), Inches(0.7), Inches(0.5))
                ntf = num_box.text_frame
                np = ntf.paragraphs[0]
                np.text = f"0{idx + 1}" if idx < 9 else f"{idx + 1}"
                np.font.name = "Arial"
                np.font.size = Pt(16)
                np.font.bold = True
                np.font.color.rgb = COLOR_BLUE

                # Text Content Box
                cbox = slide.shapes.add_textbox(Inches(1.6), Inches(current_y - 0.05), Inches(10.8), Inches(spacing - 0.15))
                ctf = cbox.text_frame
                ctf.word_wrap = True

                cp = ctf.paragraphs[0]
                cp.text = b_text
                cp.font.name = "Arial"
                cp.font.size = Pt(16)
                cp.font.color.rgb = COLOR_DARK
                cp.line_spacing = 1.25

                # Subtle hairline divider between bullets (except after the last one)
                if idx < num_bullets - 1:
                    divider = slide.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE, 
                        Inches(1.6), 
                        Inches(current_y + spacing - 0.1), 
                        Inches(10.8), 
                        Inches(0.01)
                    )
                    divider.fill.solid()
                    divider.fill.fore_color.rgb = COLOR_LINE
                    divider.line.fill.background()

        # -------------------------------------------------------------
        # FINAL SLIDE: Recommended Policy (Clean Editorial Layout)
        # -------------------------------------------------------------
        if rec_policy:
            rec_slide = prs.slides.add_slide(blank_layout)

            # Header Bar
            header_rect = rec_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
            header_rect.fill.solid()
            header_rect.fill.fore_color.rgb = COLOR_NAVY
            header_rect.line.fill.background()

            accent_bar = rec_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.04))
            accent_bar.fill.solid()
            accent_bar.fill.fore_color.rgb = COLOR_BLUE
            accent_bar.line.fill.background()

            tbox = rec_slide.shapes.add_textbox(Inches(0.9), Inches(0.2), Inches(11.5), Inches(0.75))
            ttf = tbox.text_frame
            tp = ttf.paragraphs[0]
            tp.text = "Strategic Recommendation & Policy Selection"
            tp.font.name = "Arial"
            tp.font.size = Pt(26)
            tp.font.bold = True
            tp.font.color.rgb = COLOR_WHITE

            # Map raw PDF filenames to official insurer & product names
            POLICY_DISPLAY_MAP = {
                "ABHI_Product_Brochure": "Aditya Birla Health Insurance — Activ One",
                "Care_Health_Product_Brochure": "Care Health Insurance — Care Supreme",
                "HDFC_Product_Brochure": "HDFC ERGO General Insurance — Optima Secure+",
                "Niva_Bupa_Product_Brochure": "Niva Bupa Health Insurance — ReAssure 2.0",
            }

            raw_doc = rec_policy.get('doc_name', '') if isinstance(rec_policy, dict) else str(rec_policy)
            clean_title = POLICY_DISPLAY_MAP.get(raw_doc)
            if not clean_title:
                clean_title = raw_doc.replace("_", " ").replace(".pdf", "")
            
            reason = rec_policy.get('reason', '') if isinstance(rec_policy, dict) else ''

            # Clean Content Canvas (No cards, pure typography layout)
            content_box = rec_slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.0))
            rtf = content_box.text_frame
            rtf.word_wrap = True

            rp_sub = rtf.paragraphs[0]
            rp_sub.text = "RECOMMENDED CORPORATE SOLUTION"
            rp_sub.font.name = "Arial"
            rp_sub.font.size = Pt(11)
            rp_sub.font.bold = True
            rp_sub.font.color.rgb = COLOR_BLUE
            rp_sub.space_after = Pt(8)

            rp_title = rtf.add_paragraph()
            rp_title.text = clean_title
            rp_title.font.name = "Arial"
            rp_title.font.size = Pt(32)
            rp_title.font.bold = True
            rp_title.font.color.rgb = COLOR_NAVY
            rp_title.space_after = Pt(20)

            rp_header = rtf.add_paragraph()
            rp_header.text = "Marsh Advisory Rationale & Fit:"
            rp_header.font.name = "Arial"
            rp_header.font.size = Pt(14)
            rp_header.font.bold = True
            rp_header.font.color.rgb = COLOR_DARK
            rp_header.space_after = Pt(8)

            rp_body = rtf.add_paragraph()
            rp_body.text = reason if reason else "Selected as the optimal fit matching the client's workforce demographics, claims history profile, and corporate risk parameters."
            rp_body.font.name = "Arial"
            rp_body.font.size = Pt(18)
            rp_body.font.color.rgb = COLOR_DARK
            rp_body.line_spacing = 1.3

        prs.save(output_path)
        logger.info(f"Successfully generated PowerPoint presentation at {output_path}")
        return output_path

    except Exception as e:
        error_msg = f"Failed to generate PowerPoint presentation: {str(e)}"
        logger.error(error_msg)
        return error_msg


if __name__ == "__main__":
    sample_pitch = {
        "company_name": "Syrma SGS",
        "slides": [
            {
                "slide_number": 1,
                "title": "Corporate Overview & Operational Profile",
                "bullets": [
                    {"text": "Syrma SGS is a premier electronics design and manufacturing services company with large facility operations.", "source_chunk_id": None},
                    {"text": "Key workforce risks include manufacturing floor occupational ergonomics and sedentary administrative roles.", "source_chunk_id": None}
                ]
            },
            {
                "slide_number": 2,
                "title": "Tailored Policy Benefits",
                "bullets": [
                    {"text": "Provides Day 1 cover for pre-existing chronic conditions across manufacturing personnel.", "source_chunk_id": "abhi_49"},
                    {"text": "Reimburses 100% of out-of-pocket non-medical consumables during hospitalization.", "source_chunk_id": "abhi_6"}
                ]
            }
        ],
        "recommended_policy": {
            "doc_name": "ABHI Product Brochure",
            "reason": "Activ One VIP+ uniquely addresses Syrma SGS's dual workforce dynamics by offering Day 1 chronic condition protection alongside comprehensive consumable expense reimbursement.",
            "source_chunk_id": "abhi_49"
        }
    }
    res = render_pptx(sample_pitch, "test_clean_deck.pptx")
    print("Rendered:", res)
